from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from src.common import ROOT, load_config, reporting_parameters

RESULTS = ROOT / "outputs" / "results"
TABLES = ROOT / "outputs" / "tables"

DISPLAY_COLUMNS = {
    "scenario": "Scenario",
    "normalized_objective_per_resident": "Normalized objective\nper resident",
    "facilities_open": "Open\nsites",
    "facilities_per_100k_residents": "Sites per\n100k",
    "mobile_hubs_open": "Mobile\nhubs",
    "relocated_share": "Relocated\nshare",
    "mobile_share": "Mobile\nshare",
    "digital_share": "Digital\nshare",
    "cross_boundary_share": "Cross-\nboundary",
    "mean_resident_travel_minutes": "Mean travel\n(min)",
    "p95_resident_travel_minutes": "P95 travel\n(min)",
    "resident_travel_gini": "Travel\nGini",
    "mfg_violation_share": "MFG\nviolation",
    "population_decline": "Population\ndecline",
    "rf_relocation_below_5_percent": "RF below 5%\nrelocation",
    "normalized_objective_at_minimum_rf": "Normalized objective\nat minimum RF",
    "normalized_objective_at_maximum_rf": "Normalized objective\nat maximum RF",
    "digital_share_at_maximum_rf": "Digital share at\nmaximum RF",
    "spearman_rho": "Spearman\nrho",
    "absolute_spearman_rho": "Absolute\nrho",
    "leading_driver_holdout_consistency": "Top-driver\nholdout stability",
    "capacity_multiplier": "Capacity\nmultiplier",
    "maximum_capacity_utilization_open": "Maximum open-site\nutilization",
    "sites_at_least_95_percent_capacity": "Sites ≥95%\nutilized",
    "solver_success": "Solver\nsuccess",
    "solver_mip_gap": "Solver\nMIP gap",
}

TABLE_TITLES = {
    "T1_data_and_assumptions": "Empirical inputs and model assumptions",
    "T2_scenario_design": "Scenario design",
    "T3_primary_results": "Central primary outcomes",
    "T4_capacity_robustness": "Capacity-assumption robustness",
    "T5_rf_transitions": "Relocation-friction grid diagnostics",
    "T6_sensitivity_rankings": "Designed global sensitivity rankings",
}

TABLE_NOTES = {
    "T1_data_and_assumptions": (
        "Status distinguishes empirical/public inputs, derived inputs, assumptions, "
        "and solver settings."
    ),
    "T3_primary_results": (
        "The objective is normalized model output, not expenditure or a fiscal estimate."
    ),
    "T4_capacity_robustness": (
        "Capacity multipliers vary modeled capacity assumptions, not observed staffing, "
        "throughput, service compatibility, or fiscal capacity."
    ),
    "T5_rf_transitions": (
        "RF values are model-grid diagnostics, not estimated behavioral thresholds."
    ),
    "T6_sensitivity_rankings": (
        "Spearman correlations describe importance across 96 designed "
        "Latin-hypercube ranges; they are not probabilistic uncertainty estimates."
    ),
}


def display_column(column: str) -> str:
    return DISPLAY_COLUMNS.get(column, column.replace("_", " ").title())


def display_value(value: object, column: str | None = None) -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, (float, np.floating)):
        if column in {"facilities_open", "mobile_hubs_open"}:
            return f"{value:.0f}"
        if float(value).is_integer():
            return f"{value:.0f}"
        return f"{value:.4f}"
    return str(value)


def table_1(config: dict) -> pd.DataFrame:
    summary = json.loads(
        (ROOT / "data" / "processed" / "preprocessing_summary.json").read_text(
            encoding="utf-8"
        )
    )
    values = [
        (
            "Study prefectures",
            "Empirical scope",
            "count",
            len(config["project"]["study_prefectures"]),
        ),
        ("Municipalities", "Empirical input", "count", summary["municipalities"]),
        (
            "Populated raster cells",
            "Empirical input",
            "count",
            summary["populated_cells"],
        ),
        (
            "Demand clusters",
            "Derived empirical input",
            "count",
            summary["demand_clusters"],
        ),
        (
            "Population represented",
            "Empirical input",
            "persons",
            summary["population_from_cells"],
        ),
        (
            "Medical institutions",
            "Empirical input",
            "count",
            summary["all_medical_institutions"],
        ),
        (
            "Candidate fixed assets and hubs",
            "Observed assets plus modeled hubs",
            "count",
            summary["candidate_fixed_assets_and_hubs"],
        ),
        (
            "Travel circuity multiplier",
            "Model assumption",
            "multiplier",
            config["geography"]["assumed_road_circuity"],
        ),
        (
            "Average travel speed",
            "Model assumption",
            "km/h",
            config["geography"]["assumed_average_speed_kmh"],
        ),
        (
            "Healthcare access threshold",
            "Model assumption",
            "minutes",
            config["mfg"]["healthcare_max_minutes"],
        ),
        (
            "Minimum effective provision",
            "Normative model assumption",
            "share",
            config["mfg"]["minimum_effective_provision"],
        ),
        (
            "Relative MIP-gap tolerance",
            "Solver setting",
            "relative gap",
            config["solver"]["mip_relative_gap"],
        ),
    ]
    return pd.DataFrame(values, columns=["item", "status", "unit", "value"])


def table_2() -> pd.DataFrame:
    return pd.DataFrame(
        [
            (
                "A",
                "Asset preservation",
                "Existing sites fixed open",
                "Observed regional access",
                "No",
            ),
            (
                "B",
                "Settlement consolidation",
                "Sites may close; municipal hub may open",
                "Municipal",
                "Relocation",
            ),
            (
                "C",
                "Functional provision",
                "Sites may close; mobile hub may open",
                "Municipal",
                "Relocation, mobile, digital",
            ),
            (
                "D",
                "Functional-region optimization",
                "Sites may close; mobile hub may open",
                "Cross-municipal",
                "Relocation, mobile, digital",
            ),
        ],
        columns=[
            "scenario",
            "interpretation",
            "asset_decision",
            "catchment",
            "enabled_substitution",
        ],
    )


def table_3(primary: pd.DataFrame, config: dict) -> pd.DataFrame:
    reporting = reporting_parameters(config)
    selected = primary[
        (primary["population_decline"] == reporting["baseline_decline"])
        & (primary["relocation_friction"] == reporting["central_rf"])
    ].copy()
    selected["facilities_per_100k_residents"] = (
        selected["facilities_open"] / selected["population"] * 100_000
    )
    columns = [
        "scenario",
        "cost_per_resident",
        "facilities_open",
        "facilities_per_100k_residents",
        "mobile_hubs_open",
        "relocated_share",
        "mobile_share",
        "digital_share",
        "cross_boundary_share",
        "mean_access_minutes",
        "p95_access_minutes",
        "access_gini",
        "mfg_violation_share",
    ]
    frame = selected[columns].copy()
    return frame.rename(
        columns={
            "cost_per_resident": "normalized_objective_per_resident",
            "mean_access_minutes": "mean_resident_travel_minutes",
            "p95_access_minutes": "p95_resident_travel_minutes",
            "access_gini": "resident_travel_gini",
        }
    ).sort_values("scenario")


def transition_threshold(frame: pd.DataFrame) -> float:
    below = frame[frame["relocated_share"] < 0.05]
    return float(below["relocation_friction"].min()) if len(below) else np.nan


def table_4(primary: pd.DataFrame) -> pd.DataFrame:
    records = []
    for (scenario, decline), frame in primary[
        primary["scenario"].isin(["B", "C", "D"])
    ].groupby(["scenario", "population_decline"]):
        frame = frame.sort_values("relocation_friction")
        minimum_rf = frame["relocation_friction"].min()
        maximum_rf = frame["relocation_friction"].max()
        records.append(
            {
                "scenario": scenario,
                "population_decline": decline,
                "rf_relocation_below_5_percent": transition_threshold(frame),
                "normalized_objective_at_minimum_rf": float(
                    frame.loc[
                        frame["relocation_friction"] == minimum_rf,
                        "cost_per_resident",
                    ].iloc[0]
                ),
                "normalized_objective_at_maximum_rf": float(
                    frame.loc[
                        frame["relocation_friction"] == maximum_rf,
                        "cost_per_resident",
                    ].iloc[0]
                ),
                "digital_share_at_maximum_rf": float(
                    frame.loc[
                        frame["relocation_friction"] == maximum_rf,
                        "digital_share",
                    ].iloc[0]
                ),
            }
        )
    return pd.DataFrame(records)


def table_5(
    stability: pd.DataFrame,
    robustness_summary: dict,
) -> pd.DataFrame:
    frame = stability[
        (stability["sample"] == "full") & (stability["rank"] <= 3)
    ].copy()
    consistency = {
        outcome: values["top_driver_holdout_consistency"]
        for outcome, values in robustness_summary[
            "sensitivity_rank_stability"
        ].items()
    }
    frame["leading_driver_holdout_consistency"] = np.where(
        frame["rank"] == 1,
        frame["outcome"].map(consistency),
        np.nan,
    )
    return frame[
        [
            "outcome",
            "rank",
            "parameter",
            "spearman_rho",
            "absolute_spearman_rho",
            "leading_driver_holdout_consistency",
        ]
    ]


def table_6(capacity: pd.DataFrame) -> pd.DataFrame:
    frame = capacity.rename(
        columns={"cost_per_resident": "normalized_objective_per_resident"}
    )
    return frame[
        [
            "capacity_multiplier",
            "normalized_objective_per_resident",
            "facilities_open",
            "relocated_share",
            "mobile_share",
            "digital_share",
            "cross_boundary_share",
            "mfg_violation_share",
            "maximum_capacity_utilization_open",
            "sites_at_least_95_percent_capacity",
            "solver_success",
            "solver_mip_gap",
        ]
    ].copy()


def write_editable_docx(tables: dict[str, pd.DataFrame]) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (
        section.page_height,
        section.page_width,
    )
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    document.add_heading("Editable manuscript tables", 0)
    for number, (name, frame) in enumerate(tables.items(), start=1):
        if number > 1:
            document.add_page_break()
        document.add_heading(f"Table {number}. {TABLE_TITLES[name]}", level=1)
        if note := TABLE_NOTES.get(name):
            paragraph = document.add_paragraph(note)
            paragraph.style = document.styles["Caption"]
        table = document.add_table(rows=1, cols=len(frame.columns))
        table.style = "Table Grid"
        table.autofit = False
        column_width = Inches(9.8 / len(frame.columns))
        for index, column in enumerate(frame.columns):
            table.rows[0].cells[index].text = display_column(column)
        for row in frame.itertuples(index=False):
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = display_value(value, frame.columns[index])
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.width = column_width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(7 if len(frame.columns) > 8 else 8)
                        run.bold = row_index == 0
    document.save(TABLES / "tables_editable.docx")


def format_editable_xlsx(path: Path) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for column_cells in worksheet.columns:
            maximum = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            width = min(max(maximum + 2, 10), 28)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")
    workbook.save(path)


def main() -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    config = load_config()
    primary = pd.read_csv(RESULTS / "primary_results.csv")
    stability = pd.read_csv(RESULTS / "sensitivity_rank_stability.csv")
    capacity = pd.read_csv(RESULTS / "capacity_robustness.csv")
    robustness_summary = json.loads(
        (RESULTS / "robustness_summary.json").read_text(encoding="utf-8")
    )
    tables = {
        "T1_data_and_assumptions": table_1(config),
        "T2_scenario_design": table_2(),
        "T3_primary_results": table_3(primary, config),
        "T4_capacity_robustness": table_6(capacity),
        "T5_rf_transitions": table_4(primary),
        "T6_sensitivity_rankings": table_5(stability, robustness_summary),
    }
    xlsx_path = TABLES / "tables_editable.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        for name, frame in tables.items():
            frame.to_csv(TABLES / f"{name}.csv", index=False)
            frame.to_excel(writer, sheet_name=name[:31], index=False)
    format_editable_xlsx(xlsx_path)
    write_editable_docx(tables)
    manifest = {
        name: {
            "title": TABLE_TITLES[name],
            "note": TABLE_NOTES.get(name),
            "rows": len(frame),
            "columns": list(frame.columns),
        }
        for name, frame in tables.items()
    }
    (TABLES / "table_manifest.json").write_text(
        json.dumps(manifest, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
